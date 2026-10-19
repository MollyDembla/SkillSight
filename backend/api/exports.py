"""CSV / XLSX / JSON exports, including the scheme target-setting sheet."""
from __future__ import annotations
import csv, io, json
from datetime import date
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .. import services
from ..db import rows

NOTE = "Synthetic demonstration data - not official government statistics. Planner review required before use in scheme targets."
HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")


def _csv(data: list[dict], fields: list[str], filename: str):
    s = io.StringIO()
    w = csv.DictWriter(s, fieldnames=fields, extrasaction="ignore")
    w.writeheader()
    w.writerows(data)
    return StreamingResponse(iter([s.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}"})


def _sheet(ws, data: list[dict], fields: list[str]):
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in data:
        ws.append([r.get(f) if not isinstance(r.get(f), (list, dict)) else json.dumps(r.get(f)) for f in fields])
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(10, len(f) + 2))
    ws.freeze_panes = "A2"


def _xlsx(sheets: list[tuple[str, list[dict], list[str]]], filename: str):
    wb = Workbook()
    wb.remove(wb.active)
    for name, data, fields in sheets:
        _sheet(wb.create_sheet(name[:31]), data, fields)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


FIELDS = {
    "gaps": ["state", "district", "lgd", "sector", "trade", "trade_code", "qp_code", "demand_p10", "demand_p50", "demand_p90", "demand_last12", "seats", "local_certified",
             "outflow", "inflow", "seekers", "supply", "gap", "gap_ratio", "status", "severity", "p_shortage", "p_surplus", "lead_quarter", "trust", "persistence",
             "wage_growth", "placement_rate", "invest_effect"],
    "alerts": ["id", "state", "district", "trade", "sector", "kind", "probability", "lead_quarter", "severity", "status", "assignee", "message", "updated_at"],
    "rankings": ["state", "trade_code", "title", "sector", "demand", "supply", "gap", "gap_ratio", "status", "severity", "growth", "seats"],
    "forecasts": ["node", "trade_code", "period", "p10", "p50", "p90", "invest_effect"],
    "uldi": ["period", "lgd", "trade_code", "openings", "demand_index", "n_sources", "trust"],
}


def dataset(c, name: str, fmt: str, state=None, sector=None):
    if name == "gaps":
        data = services.gap_rows(c, state=state, sector=sector, limit=10000)
    elif name == "alerts":
        data = services.alerts(c, state=state, sector=sector, limit=10000)
    elif name == "rankings":
        data = sorted(services.rollup(c, "state_trade", state=state, sector=sector), key=lambda r: r["gap_ratio"])
    elif name == "forecasts":
        data = rows(c, "SELECT f.* FROM forecasts f JOIN trades t ON t.code=f.trade_code WHERE (f.node NOT GLOB '[0-9]*' OR ?)" + (" AND t.sector=?" if sector else "") + " ORDER BY node, trade_code, period",
                    [1 if state is None else 0] + ([sector] if sector else []))
        if state:
            ids = {r["lgd"] for r in rows(c, "SELECT lgd FROM districts WHERE state=?", (state,))}
            data = [r for r in data if r["node"] in ids or r["node"] == f"ST:{state}"]
    else:
        data = rows(c, "SELECT u.* FROM uldi u JOIN districts d USING(lgd) JOIN trades t ON t.code=u.trade_code WHERE 1=1" + (" AND d.state=?" if state else "") + (" AND t.sector=?" if sector else ""),
                    [x for x in (state, sector) if x])
    fname = f"skillsight-{name}{'-' + state.replace(' ', '_') if state else ''}"
    if fmt == "json":
        return JSONResponse({"data_label": NOTE, "dataset": name, "count": len(data), "rows": data})
    if fmt == "csv":
        return _csv(data, FIELDS[name], fname + ".csv")
    return _xlsx([(name, data, FIELDS[name]), ("About", [{"note": NOTE, "generated": date.today().isoformat()}], ["note", "generated"])], fname + ".xlsx")


TARGET_FIELDS = ["state", "district", "lgd", "aspirational", "sector", "trade", "trade_code", "qp_code", "current_seats", "recommended_seats", "change",
                 "forecast_demand", "supply_before", "supply_after", "gap_ratio_before", "gap_ratio_after", "cost_change"]


def targets(res: dict, fmt: str):
    fname = f"target-sheet-{res['cycle'].replace(' ', '')}-{res['scope'].replace(' ', '_')}"
    if fmt == "json":
        return JSONResponse({"data_label": NOTE, **res})
    if fmt == "csv":
        return _csv(res["rows"], TARGET_FIELDS, fname + ".csv")
    summary = [{"metric": k, "value": v} for k, v in res["summary"].items()] + [{"metric": "cycle", "value": res["cycle"]}, {"metric": "demand window", "value": res["demand_window"]},
                                                                                   {"metric": "scope", "value": res["scope"]}, {"metric": "note", "value": NOTE}]
    assumptions = [{"assumption": k, "value": str(v)} for k, v in res["assumptions"].items()]
    by_district = {}
    for r in res["rows"]:
        d = by_district.setdefault(r["lgd"], {"state": r["state"], "district": r["district"], "current_seats": 0, "recommended_seats": 0, "cost_change": 0})
        d["current_seats"] += r["current_seats"]
        d["recommended_seats"] += r["recommended_seats"]
        d["cost_change"] += r["cost_change"]
    return _xlsx([("Summary", summary, ["metric", "value"]), ("Targets by district-trade", res["rows"], TARGET_FIELDS),
                  ("District totals", list(by_district.values()), ["state", "district", "current_seats", "recommended_seats", "cost_change"]),
                  ("Assumptions", assumptions, ["assumption", "value"])], fname + ".xlsx")
