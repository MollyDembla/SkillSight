import io
import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture(scope="module")
def client(tmp_path_factory):
 db=tmp_path_factory.mktemp("lmis")/"test.db"
 mp=pytest.MonkeyPatch()
 mp.setenv("SKILLSIGHT_DB",str(db));mp.setenv("SKILLSIGHT_REFRESH_HOURS","0");mp.setenv("SKILLSIGHT_DISABLE_LLM","1");mp.delenv("FIREBASE_PROJECT_ID",raising=False)
 from backend.pipeline import run
 run.STATE.clear()
 from backend import app as app_module
 with TestClient(app_module.app) as c:yield c
 run.STATE.clear();mp.undo()


def test_seed_is_deterministic(tmp_path,monkeypatch):
 from backend.synth import generate
 sums=[]
 for k in range(2):
  monkeypatch.setenv("SKILLSIGHT_DB",str(tmp_path/f"d{k}.db"))
  generate.seed()
  from backend.db import connect
  c=connect();sums.append((c.execute("SELECT SUM(value) FROM signals").fetchone()[0],c.execute("SELECT COUNT(*) FROM postings").fetchone()[0]));c.close()
 assert sums[0]==sums[1]


def test_reset_refuses_when_real_data_present(tmp_path,monkeypatch):
 from backend.synth import generate
 from backend.db import connect
 monkeypatch.setenv("SKILLSIGHT_DB",str(tmp_path/"r.db"));generate.seed()
 c=connect();c.execute("INSERT INTO sources(name,source_type,access_method,is_synthetic) VALUES('real','imported','csv',0)");c.commit();c.close()
 with pytest.raises(RuntimeError):generate.seed(reset=True)


def test_pipeline_quality_metrics(client):
 run=client.get('/api/v1/pipeline/runs').json()[0]
 assert run['status']=='succeeded'
 s=run['summary']
 v=s['uldi']['validation']
 # fused index must beat every single source against ground truth
 assert v['uldi_wape']<min(v['portal_only_wape'],v['ncs_only_wape'],v['epfo_only_wape'])
 for sector,w in s['uldi']['weights'].items():assert abs(sum(w.values())-1)<0.01
 f=s['forecast']
 assert f['holdout']['ensemble']<f['holdout']['linear_baseline']
 assert f['reconciliation_max_incoherence']<1e-6
 assert 0.6<f['interval_coverage_p10_p90']<0.95
 assert s['nco_coding']['accuracy']>0.9 and s['dedup']['duplicate_share']>0.1


def test_hierarchy_is_coherent(client):
 from backend.db import connect
 c=connect()
 nat=c.execute("SELECT SUM(p50) FROM forecasts WHERE node='IN' AND period='2027-03'").fetchone()[0]
 dist=c.execute("SELECT SUM(p50) FROM forecasts WHERE node GLOB '[0-9]*' AND period='2027-03'").fetchone()[0]
 sts=c.execute("SELECT SUM(p50) FROM forecasts WHERE node LIKE 'ST:%' AND period='2027-03'").fetchone()[0]
 c.close()
 assert abs(nat-dist)/nat<1e-6 and abs(nat-sts)/nat<1e-6


def test_planted_storylines_are_detected(client):
 # data entry saturation is widespread; EV battery technicians are short in Pune; Koraput hospital raises an early warning
 deo=client.get('/api/v1/rankings',params={'by':'trade','order':'surplus','limit':3}).json()
 assert 'Data Entry Operator' in [r['title'] for r in deo]
 pune=client.get('/api/v1/gaps',params={'lgd':'521','trade':'7231.0501'}).json()[0]
 assert pune['status'] in ('Acute shortage','Shortage') and pune['p_shortage']>0.8
 kor=client.get('/api/v1/alerts',params={'lgd':'398','kind':'emerging_shortage'}).json()
 assert any(a['trade']=='General Duty Assistant' for a in kor)


def test_optimizer_respects_constraints(client):
 r=client.post('/api/v1/targets/optimize',params={'rows_limit':5000},json={'state':'Odisha','max_change':0.3}).json()
 s=r['summary']
 assert s['recommended_cost']<=s['budget']+1
 assert s['unmet_demand_after']<s['unmet_demand_before']
 assert s['women_share_after']>=s['women_share_before']-1e-3
 for row in r['rows']:
  assert isinstance(row['recommended_seats'],int)
  assert row['recommended_seats']>=int(row['current_seats']*0.7)-1
 asp={}
 for row in r['rows']:
  if row['aspirational']:
   a=asp.setdefault(row['lgd'],[0,0]);a[0]+=row['current_seats'];a[1]+=row['recommended_seats']
 assert all(after>=before for before,after in asp.values())
 bad=client.post('/api/v1/targets/optimize',json={'state':'Odisha','women_target':0.8})
 assert bad.status_code==422


def test_target_sheet_export(client):
 r=client.post('/api/v1/targets/export',params={'format':'xlsx'},json={'state':'Karnataka'})
 assert r.status_code==200
 wb=load_workbook(io.BytesIO(r.content))
 assert {'Summary','Targets by district-trade','District totals','Assumptions'}<=set(wb.sheetnames)


def test_what_if_increases_demand(client):
 r=client.post('/api/v1/simulate/investment',json={'lgd':'395','sector':'RE','jobs_ops':2000,'jobs_build':3000,'ops_start':'2027-03'}).json()
 assert r['trades'] and all(t['demand_after']>=t['demand_before'] for t in r['trades'])
 assert r['total_extra_seats']>0
 assert client.post('/api/v1/simulate/investment',json={'lgd':'nope','sector':'RE','jobs_ops':20,'ops_start':'2027-03'}).status_code==422


def test_pathways_curriculum_siting(client):
 p=client.get('/api/v1/pathways').json()
 assert p['total_absorbable']>0 and all(x['skill_overlap']>=0.25 for x in p['pathways'])
 cur=client.get('/api/v1/curriculum').json()
 ev=next(x for x in cur if x['trade_code']=='7231.0501')
 assert ev['missing_skills'] and ev['freshness_score']<75
 assert client.get('/api/v1/siting',params={'limit':5}).json()[0]['priority_score']>0


def test_alert_workflow_and_digest(client):
 a=client.get('/api/v1/alerts',params={'limit':1}).json()[0]
 r=client.patch(f"/api/v1/alerts/{a['id']}",json={'status':'assigned','assignee':'Odisha SSDM'}).json()
 assert r['status']=='assigned'
 assert 'body' in client.get('/api/v1/alerts/digest',params={'channel':'whatsapp'}).json()


def test_copilot_offline_answers_with_data(client):
 r=client.post('/api/v1/copilot/chat',json={'message':'Where are the biggest shortages in Odisha?'}).json()
 assert r['mode']=='offline' and 'Odisha' in r['answer'] and r['data']['rows']
 r=client.post('/api/v1/copilot/chat',json={'message':'forecast for EV battery technician in Pune'}).json()
 assert 'EV Battery Technician' in r['answer']


def test_nco_classifier_and_import_harmonisation(client):
 r=client.post('/api/v1/nco/classify',json={'titles':['solar mistri','ward boy hospital','Helper']}).json()['results']
 assert r[0]['code']=='7411.0101' and r[1]['code']=='5321.0100' and r[2]['needs_review']
 good='period,state,district,occupation,value\n2026-09,Karnataka,Bangalore,Solar technician,40\n'
 res=client.post('/api/v1/import',files={'file':('a.csv',good,'text/csv')}).json()
 assert res['rows_imported']==1 and res['harmonisation'][0]['district']=='Bengaluru Urban'
 bad='period,state,district,occupation,value\n2026-09,Karnataka,Atlantis,Solar technician,40\n'
 assert client.post('/api/v1/import',files={'file':('b.csv',bad,'text/csv')}).status_code==422


def test_exports_reports_and_keys(client):
 for ds in ['gaps','alerts','rankings']:
  assert client.get(f'/api/v1/exports/{ds}.csv',params={'state':'Odisha'}).status_code==200
 assert client.get('/api/v1/exports/forecasts.json').json()['count']>0
 html=client.get('/api/v1/reports/dsdp/388',params={'lang':'hi'}).text
 assert 'जिला कौशल विकास योजना' in html and 'Ganjam' in html
 k=client.post('/api/v1/api-keys',json={'name':'scheme-mis'}).json()
 assert k['key'].startswith('sk_lmis_') and client.get('/api/v1/api-keys').json()[0]['prefix']==k['key'][:12]


def test_gemini_copilot_function_calling_loop(client,monkeypatch):
 """Drive the Gemini tool loop with a fake model: one function call, then a final answer."""
 from types import SimpleNamespace
 from google.genai import types
 from backend import llm
 assert len(llm._gemini_tools()[0].function_declarations)==len(llm.TOOL_SPECS)
 calls=[]
 def fake_generate(contents,**config):
  calls.append(contents)
  if len(calls)==1:
   fc=SimpleNamespace(name='query_gaps',args={'state':'Odisha','status':'Acute shortage','limit':3})
   content=types.Content(role='model',parts=[types.Part.from_function_call(name='query_gaps',args=fc.args)])
   return SimpleNamespace(candidates=[SimpleNamespace(content=content)],function_calls=[fc],text=None)
  assert contents[-1].parts[0].function_response.response['result']  # tool output was sent back
  return SimpleNamespace(candidates=[SimpleNamespace(content=types.Content(role='model',parts=[types.Part.from_text(text='done')]))],function_calls=None,text='Odisha shortages: ...')
 monkeypatch.setattr(llm,'available',lambda:True);monkeypatch.setattr(llm,'_generate',fake_generate)
 r=client.post('/api/v1/copilot/chat',json={'message':'shortages in Odisha?'}).json()
 assert r['mode']=='gemini' and r['tools'][0]['name']=='query_gaps' and r['data']['rows'] and len(calls)==2
